"""Step 4: translate the three NBA-game net uplifts into low/base/high demand
scenarios. Same numbers for every host market."""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
ANALOGS_CSV = ROOT / "outputs" / "analogs_table.csv"
SCENARIOS_CSV = ROOT / "outputs" / "scenarios.csv"
PLAN_XLSX = ROOT / "outputs" / "plan.xlsx"

# scenario -> (source event_id, expected net_peak_uplift -- checked, not invented)
SCENARIO_SOURCE = {"low": "N2", "base": "N3", "high": "N4"}
EXPECTED_NET_PEAK = {"N2": -0.112, "N3": 0.261, "N4": 0.306}
TOLERANCE = 0.005  # 0.5 percentage points

K_VALUES = [0.25, 0.5]
PROBABILITIES = {"low": 0.40, "base": 0.40, "high": 0.20}  # judgement call, not derived from data

COUNTRIES = ["UK", "France", "Spain", "Italy", "Germany", "Turkey"]


def load_interest_uplifts() -> tuple[dict, list[tuple]]:
    """Reads net_peak_uplift (basketball, exposed) for N2/N3/N4 from
    analogs_table.csv, checks it against the expected values, and floors any
    negative value at 0%. Returns (scenario -> info dict, floor log)."""
    if not ANALOGS_CSV.exists():
        sys.exit(f"STOP: {ANALOGS_CSV} not found -- run src/analogs.py first.")
    df = pd.read_csv(ANALOGS_CSV)
    df = df[(df["topic"] == "basketball") & (df["exposed"])]

    result = {}
    floor_log = []
    for scenario, event_id in SCENARIO_SOURCE.items():
        row = df[df["event_id"] == event_id]
        if row.empty:
            sys.exit(f"STOP: no basketball/exposed row for event {event_id} in {ANALOGS_CSV}")
        raw = float(row["net_peak_uplift"].iloc[0])
        event_name = str(row["event_name"].iloc[0])
        country = str(row["country"].iloc[0])

        expected = EXPECTED_NET_PEAK[event_id]
        if abs(raw - expected) > TOLERANCE:
            sys.exit(
                f"STOP: {event_id} ({event_name}, {country}) net_peak_uplift={raw * 100:.2f}% "
                f"differs from the expected {expected * 100:.1f}% by more than "
                f"{TOLERANCE * 100:.1f} points -- re-check inputs before trusting these scenarios."
            )

        floored = raw < 0
        interest_uplift = max(raw, 0.0)
        if floored:
            floor_log.append((scenario, event_id, event_name, country, raw))

        result[scenario] = {
            "event_id": event_id,
            "event_name": event_name,
            "country": country,
            "raw_net_peak_uplift": raw,
            "interest_uplift": interest_uplift,
            "floored": floored,
        }
    return result, floor_log


def expected_interest_uplift(interest: dict) -> float:
    return sum(interest[s]["interest_uplift"] * PROBABILITIES[s] for s in ("low", "base", "high"))


def build_scenarios_table(interest: dict) -> pd.DataFrame:
    """k, scenario, interest_uplift, demand_uplift, probability -- same numbers
    for every host market, so no country column. One "expected" row per k
    (the probability-weighted average; probability left blank since it's not
    one of the three mutually exclusive scenarios)."""
    expected_interest = expected_interest_uplift(interest)
    rows = []
    for k in K_VALUES:
        for scenario in ("low", "base", "high"):
            info = interest[scenario]
            rows.append(
                {
                    "k": k,
                    "scenario": scenario,
                    "interest_uplift": info["interest_uplift"],
                    "demand_uplift": info["interest_uplift"] * k,
                    "probability": PROBABILITIES[scenario],
                }
            )
        rows.append(
            {
                "k": k,
                "scenario": "expected",
                "interest_uplift": expected_interest,
                "demand_uplift": expected_interest * k,
                "probability": "",
            }
        )
    return pd.DataFrame(rows, columns=["k", "scenario", "interest_uplift", "demand_uplift", "probability"])


# ---------------------------------------------------------------------------
# Excel output -- one sheet, "Scenarios": one small table per k.
# ---------------------------------------------------------------------------

PCT_FORMAT = "0.0%"
HEADER_FONT = Font(bold=True)


def write_scenarios_sheet(ws, interest: dict) -> None:
    ws.title = "Scenarios"
    ws["A1"] = "Demand uplift by scenario (basketball category, vs. baseline plan)"
    ws["A1"].font = Font(bold=True, size=12)

    expected_interest = expected_interest_uplift(interest)
    row = 3
    for k in K_VALUES:
        ws.cell(row=row, column=1, value=f"k = {k}").font = HEADER_FONT
        row += 1
        for col, h in enumerate(["Low", "Base", "High", "Expected"], start=1):
            ws.cell(row=row, column=col, value=h).font = HEADER_FONT
        row += 1
        for col, scenario in zip((1, 2, 3), ("low", "base", "high")):
            cell = ws.cell(row=row, column=col, value=interest[scenario]["interest_uplift"] * k)
            cell.number_format = PCT_FORMAT
        cell = ws.cell(row=row, column=4, value=expected_interest * k)
        cell.number_format = PCT_FORMAT
        row += 1
        ws.cell(
            row=row, column=1,
            value=(
                f"k = share of extra interest assumed to convert to extra purchases. "
                f"Expected = {PROBABILITIES['low']*100:.0f}% low / {PROBABILITIES['base']*100:.0f}% base / "
                f"{PROBABILITIES['high']*100:.0f}% high (judgement call)."
            ),
        )
        row += 2  # blank row between the two k blocks

    for col, width in {"A": 42, "B": 12, "C": 12, "D": 12}.items():
        ws.column_dimensions[col].width = width


def write_plan_xlsx(interest: dict) -> None:
    wb = Workbook()
    write_scenarios_sheet(wb.active, interest)
    wb.save(PLAN_XLSX)


def print_summary(interest: dict, floor_log: list[tuple], scenarios_df: pd.DataFrame) -> None:
    print("=== Scenario inputs (net_peak_uplift, basketball, exposed) ===")
    for scenario, info in interest.items():
        note = " -> FLOORED to 0%" if info["floored"] else ""
        print(
            f"  {scenario:<5} {info['event_name']:<18} ({info['country']}): "
            f"{info['raw_net_peak_uplift'] * 100:+.2f}%{note}"
        )

    if floor_log:
        print("\nFloor applied:")
        for scenario, event_id, event_name, country, raw in floor_log:
            print(f"  {scenario} ({event_id} {event_name}, {country}): {raw * 100:+.2f}% -> 0.0%")

    expected_interest = expected_interest_uplift(interest)
    print("\n=== Scenario table (same for all 6 countries) ===")
    for k in K_VALUES:
        print(f"\nk = {k}")
        for scenario in ("low", "base", "high"):
            info = interest[scenario]
            print(
                f"  {scenario:<8} demand_uplift={info['interest_uplift'] * k * 100:+.2f}%  "
                f"(probability={PROBABILITIES[scenario] * 100:.0f}%)"
            )
        print(f"  {'expected':<8} demand_uplift={expected_interest * k * 100:+.2f}%")

    print(f"\nRows written to {SCENARIOS_CSV.relative_to(ROOT)}: {len(scenarios_df)}")


def main() -> None:
    interest, floor_log = load_interest_uplifts()
    scenarios_df = build_scenarios_table(interest)
    scenarios_df.to_csv(SCENARIOS_CSV, index=False)
    write_plan_xlsx(interest)
    print_summary(interest, floor_log, scenarios_df)


if __name__ == "__main__":
    main()
