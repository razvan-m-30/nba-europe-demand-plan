"""Step 5: newsvendor buy rule -- translate demand scenarios into a recommended
buy level, balancing missed full-price sales against markdown losses."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
SCENARIOS_CSV = ROOT / "outputs" / "scenarios.csv"
NEWSVENDOR_CSV = ROOT / "outputs" / "newsvendor.csv"
PLAN_XLSX = ROOT / "outputs" / "plan.xlsx"

# Nike gross margin, FY2025 (fiscal year ended 2025-05-31): 42.7%, down from
# 44.6% in FY2024. Verified via web search against Nike's FY2025 10-K and
# investor press release:
#   https://www.sec.gov/Archives/edgar/data/320187/000032018725000047/nke-20250531.htm
#   https://investors.nike.com/news-events/press-releases/detail/551/nike-inc-reports-fiscal-2025-fourth-quarter-and-full-year-results
GROSS_MARGIN = 0.427
GROSS_MARGIN_FY = "FY2025 (ended 2025-05-31)"
GROSS_MARGIN_SOURCE = (
    "Nike 10-K FY2025: "
    "https://www.sec.gov/Archives/edgar/data/320187/000032018725000047/nke-20250531.htm"
)

MARKDOWN_VALUES = [0.30, 0.50]
K_VALUES = [0.25, 0.5]
CUMULATIVE_PROBABILITY = {"low": 0.40, "base": 0.80, "high": 1.00}
SCENARIO_ORDER = ["low", "base", "high"]


def load_scenarios() -> pd.DataFrame:
    if not SCENARIOS_CSV.exists():
        raise SystemExit(f"STOP: {SCENARIOS_CSV} not found -- run src/scenarios.py first.")
    return pd.read_csv(SCENARIOS_CSV)


def recommended_scenario(critical_ratio: float) -> str:
    """First scenario (low -> base -> high) whose cumulative probability
    reaches the critical ratio."""
    for scenario in SCENARIO_ORDER:
        if CUMULATIVE_PROBABILITY[scenario] >= critical_ratio:
            return scenario
    return SCENARIO_ORDER[-1]


def build_newsvendor_table(scenarios: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for k in K_VALUES:
        subset = scenarios[scenarios["k"] == k]
        demand_by_scenario = {
            s: float(subset.loc[subset["scenario"] == s, "demand_uplift"].iloc[0]) for s in SCENARIO_ORDER
        }
        expected = float(subset.loc[subset["scenario"] == "expected", "demand_uplift"].iloc[0])

        for markdown in MARKDOWN_VALUES:
            cu = GROSS_MARGIN
            co = markdown
            critical_ratio = cu / (cu + co)
            scenario = recommended_scenario(critical_ratio)
            rows.append(
                {
                    "k": k,
                    "markdown": markdown,
                    "critical_ratio": critical_ratio,
                    "recommended_scenario": scenario,
                    "recommended_buy": demand_by_scenario[scenario],
                    "expected_uplift": expected,
                }
            )
    return pd.DataFrame(
        rows,
        columns=["k", "markdown", "critical_ratio", "recommended_scenario", "recommended_buy", "expected_uplift"],
    )


# ---------------------------------------------------------------------------
# Excel: add a "Buy" sheet next to "Scenarios"
# ---------------------------------------------------------------------------

PCT_FORMAT = "0.0%"
HEADER_FONT = Font(bold=True)


def write_buy_sheet(table: pd.DataFrame) -> None:
    if not PLAN_XLSX.exists():
        raise SystemExit(f"STOP: {PLAN_XLSX} not found -- run src/scenarios.py first.")
    wb = load_workbook(PLAN_XLSX)
    if "Buy" in wb.sheetnames:
        del wb["Buy"]
    ws = wb.create_sheet("Buy")

    ws["A1"] = "Buy recommendation (newsvendor: Cu = gross margin, Co = markdown depth)"
    ws["A1"].font = Font(bold=True, size=12)

    ws["A3"] = "Inputs"
    ws["A3"].font = HEADER_FONT
    ws["A4"] = "Gross margin (Cu)"
    ws["B4"] = GROSS_MARGIN
    ws["B4"].number_format = PCT_FORMAT
    ws["C4"] = f"{GROSS_MARGIN_FY} -- {GROSS_MARGIN_SOURCE}"
    ws["A5"] = "Markdown depth (Co)"
    ws["B5"] = "30% and 50%"
    ws["C5"] = "Tested at two depths; a deeper markdown lowers the critical ratio and pushes the buy down"
    ws["A6"] = "Cumulative probability"
    ws["B6"] = "low 40%, base 80%, high 100%"
    ws["C6"] = "From scenarios.csv's low/base/high probabilities (40/40/20)"

    row = 8
    ws.cell(row=row, column=1, value="Results").font = HEADER_FONT
    row += 1
    headers = ["k", "Markdown", "Critical ratio", "Recommended scenario", "Buy (% above plan)"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=row, column=col, value=h).font = HEADER_FONT
    row += 1
    for _, r in table.iterrows():
        ws.cell(row=row, column=1, value=r["k"])
        c = ws.cell(row=row, column=2, value=r["markdown"])
        c.number_format = PCT_FORMAT
        c = ws.cell(row=row, column=3, value=r["critical_ratio"])
        c.number_format = "0.000"
        ws.cell(row=row, column=4, value=r["recommended_scenario"])
        c = ws.cell(row=row, column=5, value=r["recommended_buy"])
        c.number_format = PCT_FORMAT
        row += 1

    row += 1
    ws.cell(
        row=row, column=1,
        value="Buy = the scenario level that balances missed full-price sales against markdown losses.",
    )

    widths = {"A": 24, "B": 16, "C": 70, "D": 22, "E": 20}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(PLAN_XLSX)


def print_summary(table: pd.DataFrame) -> None:
    print(f"=== Gross margin ===")
    print(f"  {GROSS_MARGIN * 100:.1f}% ({GROSS_MARGIN_FY})")
    print(f"  Source: {GROSS_MARGIN_SOURCE}")

    print("\n=== Newsvendor result table ===")
    for _, r in table.iterrows():
        print(
            f"  k={r['k']:<5} markdown={r['markdown']*100:.0f}%  "
            f"critical_ratio={r['critical_ratio']:.3f}  "
            f"recommended={r['recommended_scenario']:<5}  "
            f"buy={r['recommended_buy']*100:+.2f}%  "
            f"expected={r['expected_uplift']*100:+.2f}%"
        )


def main() -> None:
    scenarios = load_scenarios()
    table = build_newsvendor_table(scenarios)
    table.to_csv(NEWSVENDOR_CSV, index=False)
    write_buy_sheet(table)
    print_summary(table)


if __name__ == "__main__":
    main()
